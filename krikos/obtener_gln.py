import re
import unicodedata

from difflib import SequenceMatcher

try:
    import PyPDF2
except ImportError:
    import pypdf as PyPDF2

from openpyxl import load_workbook

from config import ARCHIVO_MAPEO_SUCURSALES


def _normalizar(texto):

    if texto is None:
        return ""

    texto = str(
        texto
    ).upper().strip()

    texto = unicodedata.normalize(
        "NFD",
        texto
    )

    texto = "".join(
        caracter
        for caracter in texto
        if unicodedata.category(
            caracter
        )
        != "Mn"
    )

    texto = re.sub(
        r"[^A-Z0-9]+",
        " ",
        texto
    )

    return re.sub(
        r"\s+",
        " ",
        texto
    ).strip()


def _extraer_texto_pdf(
    ruta_pdf,
):

    with open(
        ruta_pdf,
        "rb",
    ) as archivo:

        lector = PyPDF2.PdfReader(
            archivo
        )

        return "\n".join(
            (
                pagina.extract_text()
                or ""
            )
            for pagina
            in lector.pages
        )


def _extraer_datos_cliente(
    texto,
):

    compacto = " ".join(
        texto.split()
    )


    # ==========================================
    # NÚMERO DE CLIENTE
    # ==========================================

    patrones = [

        r"N[º°]?\s*Clie\.?\s*:?\s*(\d+)",

        r"Nro\.?\s*Clie\.?\s*:?\s*(\d+)",

        r"N\.?\s*Cliente\s*:?\s*(\d+)",
    ]


    numero_cliente = ""


    for patron in patrones:

        coincidencia = re.search(
            patron,
            compacto,
            re.IGNORECASE,
        )

        if coincidencia:

            numero_cliente = (
                coincidencia.group(
                    1
                )
            )

            break


    # ==========================================
    # NOMBRE
    # ==========================================

    nombre = ""


    coincidencia = re.search(
        (
            r"Sr\.?(?:\(s\))?\s*[:.]?\s*"
            r"(.+?)"
            r"(?=\s+Domicilio\b"
            r"|\s+IVA\s*:"
            r"|\s+Telefono\s*:"
            r"|\s+Tel[eé]fono\s*:)"
        ),
        compacto,
        re.IGNORECASE,
    )


    if coincidencia:

        nombre = (
            coincidencia
            .group(1)
            .strip()
        )


    # ==========================================
    # DOMICILIO
    # ==========================================

    domicilio = ""


    coincidencia = re.search(
        (
            r"Domicilio\s*:?\s*"
            r"(.+?)"
            r"(?=\s+IVA\s*:"
            r"|\s+F\.?\s*Pago\s*:"
            r"|\s+Telefono\s*:"
            r"|\s+Tel[eé]fono\s*:"
            r"|\s+Localidad\s*:)"
        ),
        compacto,
        re.IGNORECASE,
    )


    if coincidencia:

        domicilio = (
            coincidencia
            .group(1)
            .strip()
        )


    # ==========================================
    # LOCALIDAD
    # ==========================================

    localidad = ""


    coincidencia = re.search(
        (
            r"Localidad\s*:?\s*"
            r"(.+?)"
            r"(?=\s+Provincia\s*:"
            r"|\s+CUIT\s*:"
            r"|\s+I\.?Brutos\s*:"
            r"|\s+Cod\.Act\.?\s*:)"
        ),
        compacto,
        re.IGNORECASE,
    )


    if coincidencia:

        localidad = (
            coincidencia
            .group(1)
            .strip()
        )


    return {

        "nro_cliente":
            numero_cliente,

        "nombre":
            nombre,

        "domicilio":
            domicilio,

        "localidad":
            localidad,
    }


def _similitud(
    texto_a,
    texto_b,
):

    texto_a = _normalizar(
        texto_a
    )

    texto_b = _normalizar(
        texto_b
    )


    if (
        not texto_a
        or
        not texto_b
    ):

        return 0.0


    if (
        texto_a in texto_b
        or
        texto_b in texto_a
    ):

        return 1.0


    return SequenceMatcher(
        None,
        texto_a,
        texto_b,
    ).ratio()


def obtener_gln_factura(
    ruta_pdf,
):

    resultado_vacio = {

        "encontrado":
            False,

        "gln":
            "",

        "sucursal_json":
            "",

        "sucursal_trilay":
            "",

        "domicilio_json":
            "",

        "metodo":
            "",

        "confianza":
            0.0,
    }


    try:

        texto = _extraer_texto_pdf(
            ruta_pdf
        )


        cliente = _extraer_datos_cliente(
            texto
        )


        if not ARCHIVO_MAPEO_SUCURSALES.is_file():

            raise RuntimeError(
                "No existe el archivo "
                "Mapeo_Sucursales_Completo.xlsx "
                f"en {ARCHIVO_MAPEO_SUCURSALES}"
            )


        workbook = load_workbook(
            ARCHIVO_MAPEO_SUCURSALES,
            read_only=True,
            data_only=True,
        )


        if (
            "Mapeo Sucursales"
            in workbook.sheetnames
        ):

            hoja = workbook[
                "Mapeo Sucursales"
            ]

        else:

            hoja = workbook[
                workbook.sheetnames[0]
            ]


        filas = []


        for (
            sucursal_trilay,
            sucursal_json,
            domicilio_json,
            gln,
        ) in hoja.iter_rows(
            min_row=2,
            values_only=True,
        ):

            if not gln:
                continue


            filas.append(
                {

                    "sucursal_trilay":
                        str(
                            sucursal_trilay
                            or ""
                        ).strip(),

                    "sucursal_json":
                        str(
                            sucursal_json
                            or ""
                        ).strip(),

                    "domicilio_json":
                        str(
                            domicilio_json
                            or ""
                        ).strip(),

                    "gln":
                        str(
                            gln
                        ).strip(),
                }
            )


        workbook.close()


        # ==========================================
        # 1. BUSCAR POR NÚMERO DE CLIENTE
        # ==========================================

        numero_cliente = (
            cliente[
                "nro_cliente"
            ]
        )


        if numero_cliente:

            patron = re.compile(
                rf"^0*"
                rf"{re.escape(numero_cliente)}"
                rf"(?:\D|$)"
            )


            coincidencias = [

                fila
                for fila in filas

                if patron.search(
                    fila[
                        "sucursal_trilay"
                    ]
                )
            ]


            if len(
                coincidencias
            ) == 1:

                mejor = (
                    coincidencias[0]
                )


                return {

                    "encontrado":
                        True,

                    **mejor,

                    "metodo":
                        "numero_cliente",

                    "confianza":
                        1.0,
                }


            if len(
                coincidencias
            ) > 1:

                filas = coincidencias


        # ==========================================
        # 2. FALLBACK POR SIMILITUD
        # ==========================================

        mejor = None

        mejor_score = 0.0


        for fila in filas:

            score_nombre = max(

                _similitud(
                    cliente[
                        "nombre"
                    ],
                    fila[
                        "sucursal_trilay"
                    ],
                ),

                _similitud(
                    cliente[
                        "nombre"
                    ],
                    fila[
                        "sucursal_json"
                    ],
                ),
            )


            score_domicilio = (
                _similitud(
                    cliente[
                        "domicilio"
                    ],
                    fila[
                        "domicilio_json"
                    ],
                )
            )


            score_localidad = (
                _similitud(
                    cliente[
                        "localidad"
                    ],
                    fila[
                        "sucursal_json"
                    ],
                )
            )


            score = (
                score_nombre
                * 0.70
            ) + (
                score_domicilio
                * 0.20
            ) + (
                score_localidad
                * 0.10
            )


            if score > mejor_score:

                mejor_score = score

                mejor = fila


        # ==========================================
        # NIVEL MÍNIMO DE CONFIANZA
        # ==========================================

        if (
            mejor
            and
            mejor_score >= 0.58
        ):

            return {

                "encontrado":
                    True,

                **mejor,

                "metodo":
                    "similitud",

                "confianza":
                    round(
                        mejor_score,
                        3
                    ),
            }


        return resultado_vacio


    except Exception as error:

        print(
            f"Error obteniendo GLN: "
            f"{error}"
        )

        return resultado_vacio